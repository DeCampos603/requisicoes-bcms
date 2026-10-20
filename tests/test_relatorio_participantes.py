# -*- coding: utf-8 -*-
"""Testes automatizados do módulo de Relatório de Pregões Participantes."""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.relatorio_participantes import (
    carregar_catalogo_uasgs,
    carregar_pregoes_homologados,
    consolidar_dados,
    gerar_excel,
    gerar_relatorio,
)


def test_catalogo_uasgs():
    uasgs = carregar_catalogo_uasgs()
    assert isinstance(uasgs, dict)
    assert "160329" in uasgs
    assert "160238" in uasgs
    assert "160200" in uasgs


def test_carregar_pregoes_homologados():
    pregoes, itens = carregar_pregoes_homologados()
    assert isinstance(pregoes, list)
    assert isinstance(itens, list)
    assert len(pregoes) >= 30
    assert len(itens) >= 1000

    # Valida estrutura do primeiro pregão
    p0 = pregoes[0]
    assert "ugg" in p0
    assert "pregao" in p0
    assert "n_itens" in p0
    assert "valor_cota_bcms" in p0
    assert "valor_saldo_bcms" in p0
    assert "valor_global_ata" in p0
    assert p0["valor_cota_bcms"] >= 0

    # Valida estrutura do primeiro item
    it0 = itens[0]
    assert "item" in it0
    assert "catmat" in it0
    assert "descricao" in it0
    assert "valor_unitario" in it0
    assert "qtd_cota_bcms" in it0
    assert "qtd_saldo_bcms" in it0
    assert it0["valor_unitario"] >= 0


def test_consolidar_dados():
    dados = consolidar_dados(atualizar_pncp=False)
    assert "kpis" in dados
    assert "pregoes_homologados" in dados
    assert "itens_detalhados" in dados
    assert "pncp_em_andamento" in dados
    assert "tabela_oms" in dados

    kpis = dados["kpis"]
    assert kpis["total_pregoes_homologados"] >= 30
    assert kpis["total_itens_bcms"] >= 1000
    assert kpis["valor_cota_geral"] > 0
    assert kpis["valor_saldo_geral"] > 0
    assert 0 <= kpis["pct_saldo_consumido"] <= 100


def test_gerar_excel_completo():
    with tempfile.TemporaryDirectory() as tmpdir:
        destino = Path(tmpdir) / "teste_relatorio.xlsx"
        dados = consolidar_dados(atualizar_pncp=False)
        gerado = gerar_excel(dados, caminho_saida=destino)

        assert gerado.is_file()
        assert gerado.stat().st_size > 50_000

        # Inspeciona com openpyxl
        wb = openpyxl.load_workbook(gerado, data_only=False)
        assert len(wb.sheetnames) == 5
        assert "Painel Resumo & KPIs" in wb.sheetnames
        assert "Pregões em Andamento (PNCP)" in wb.sheetnames
        assert "Pregões Homologados (Atas)" in wb.sheetnames
        assert "Itens Detalhados (BCMS)" in wb.sheetnames
        assert "Cruzamento & Alertas" in wb.sheetnames

        # Confere cabeçalhos e valores
        ws1 = wb["Painel Resumo & KPIs"]
        assert "BATALHÃO CENTRAL DE MANUTENÇÃO E SUPRIMENTO" in str(ws1["A1"].value)

        ws3 = wb["Pregões Homologados (Atas)"]
        assert ws3.max_row >= 30
        assert ws3.cell(1, 1).value == "UASG Gerenciadora"

        ws4 = wb["Itens Detalhados (BCMS)"]
        assert ws4.max_row >= 1000

        wb.close()


def test_api_relatorio_resumo():
    cliente = TestClient(app, headers={"Host": "127.0.0.1"})
    resp = cliente.get("/api/relatorio/participantes/resumo")
    assert resp.status_code == 200
    data = resp.json()
    assert "kpis" in data
    assert data["kpis"]["total_pregoes_homologados"] >= 30


def test_api_relatorio_excel():
    cliente = TestClient(app, headers={"Host": "127.0.0.1"})
    resp = cliente.get("/api/relatorio/participantes/excel")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert len(resp.content) > 50_000
