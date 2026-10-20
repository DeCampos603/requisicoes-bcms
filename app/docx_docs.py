# -*- coding: utf-8 -*-
"""Geração de documentos em formato Microsoft Word (.docx) no padrão dos DIEx do BCMS.

Segue a mesma estrutura e diagramação oficial dos DIEx em PDF gerados por app.pdf_docs:
  * Papel A4 com margens oficiais: Esquerda 2,0 cm, Direita 1,27 cm, Superior 1,27 cm, Inferior 1,27 cm;
  * Fonte institucional Times New Roman 12 pt;
  * Brasão das Armas Nacionais e cabeçalho institucional centralizado em negrito;
  * Bloco de identificação (DIEx/NUP à esquerda, local e data à direita; Do / Ao / Assunto / Anexos);
  * Parágrafos justificados com suporte a negrito e formatação;
  * Quadro de condições da despesa (quando aplicável);
  * Tabelas de itens com grade, cabeçalho sombreado e alinhamentos adequados;
  * Bloco de assinatura do requisitante;
  * Quadro de despacho com duas colunas (Fiscal Administrativo e Ordenador de Despesas).
"""

from __future__ import annotations

import html
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

from . import config

_RE_TAG = re.compile(r"(</?[a-zA-Z0-9_-]+[^>]*>)")


def _brasao_bytes() -> bytes | None:
    arq = config.ASSETS / "brasao_armas.png"
    if not arq.is_file():
        return None
    try:
        im = Image.open(arq).convert("RGBA")
        caixa = im.getchannel("A").getbbox()
        if caixa:
            im = im.crop(caixa)
        fundo = Image.new("RGB", im.size, (255, 255, 255))
        fundo.paste(im, mask=im.getchannel("A"))
        fundo.thumbnail((420, 420))
        buf = io.BytesIO()
        fundo.save(buf, "JPEG", quality=92, optimize=True)
        return buf.getvalue()
    except Exception:
        return None


def _limpar_texto(txt: str) -> str:
    if not txt:
        return ""
    s = html.unescape(txt)
    return s.replace("&nbsp;", " ").replace("\u00a0", " ")


def _aplicar_formatacao_paragrafo(p, texto_com_tags: str, tamanho_pt: float = 12.0,
                                  negrito_padrao: bool = False, cor_rgb: RGBColor | None = None,
                                  italico: bool = False, sublinhado: bool = False) -> None:
    tokens = _RE_TAG.split(texto_com_tags or "")
    negrito = negrito_padrao
    italico_atual = italico
    sublinhado_atual = sublinhado

    for tok in tokens:
        if not tok:
            continue
        tok_low = tok.lower().strip()
        if tok_low == "<b>":
            negrito = True
        elif tok_low == "</b>":
            negrito = negrito_padrao
        elif tok_low == "<i>":
            italico_atual = True
        elif tok_low == "</i>":
            italico_atual = italico
        elif tok_low == "<u>":
            sublinhado_atual = True
        elif tok_low == "</u>":
            sublinhado_atual = sublinhado
        elif tok_low.startswith("<") and tok_low.endswith(">"):
            continue
        else:
            texto_limpo = _limpar_texto(tok)
            run = p.add_run(texto_limpo)
            run.font.name = "Times New Roman"
            run.font.size = Pt(tamanho_pt)
            run.bold = negrito
            run.italic = italico_atual
            run.underline = sublinhado_atual
            if cor_rgb:
                run.font.color.rgb = cor_rgb


def _definir_bordas_celula(cell, top: str = "single", bottom: str = "single",
                           left: str = "single", right: str = "single",
                           cor: str = "333333", sz: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{cor}"/>\n'
        f'  <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{cor}"/>\n'
        f'  <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{cor}"/>\n'
        f'  <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{cor}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def _sombrear_celula(cell, cor_hex: str = "F0F0F0") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{cor_hex}"/>')
    tcPr.append(shd)


def _margens_celula(cell, top: int = 100, bottom: int = 100, left: int = 120, right: int = 120) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _adicionar_campo_pagina_docx(paragraph, tamanho_pt: float = 9.0, cor_rgb: RGBColor | None = None) -> None:
    """Adiciona campos dinâmicos PAGE e NUMPAGES para 'Página X de Y' no Word."""
    r1 = paragraph.add_run("Página ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(tamanho_pt)
    r1.bold = True
    if cor_rgb:
        r1.font.color.rgb = cor_rgb

    fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls("w"))
    paragraph._p.append(fld_page)

    r2 = paragraph.add_run(" de ")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(tamanho_pt)
    r2.bold = True
    if cor_rgb:
        r2.font.color.rgb = cor_rgb

    fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls("w"))
    paragraph._p.append(fld_numpages)


def gerar_docx_documento(doc: dict, destino: Path, anexos: list[tuple[str, Path]] | None = None) -> Path:
    """Gera um documento .docx a partir da estrutura doc montada por app.textos, incluindo anexos se fornecidos."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    d = Document()

    # Configuração de Página A4 e Margens (Esquerda 2,0 cm; Superior/Direita/Inferior 1,27 cm)
    for section in d.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.27)
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.different_first_page_header_footer = False

        # Rodapé oficial com informações do DIEx/NUP e ordem das páginas
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.text = ""
        p_ft.paragraph_format.space_before = Pt(0)
        p_ft.paragraph_format.space_after = Pt(0)

        tbl_ft = footer.add_table(rows=1, cols=2, width=Cm(17.73))
        tbl_ft.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_ft.autofit = False
        tbl_ft.columns[0].width = Cm(13.5)
        tbl_ft.columns[1].width = Cm(4.23)
        _definir_bordas_celula(tbl_ft.cell(0, 0), top="single", bottom="none", left="none", right="none", cor="888888", sz="4")
        _definir_bordas_celula(tbl_ft.cell(0, 1), top="single", bottom="none", left="none", right="none", cor="888888", sz="4")
        _margens_celula(tbl_ft.cell(0, 0), top=60, bottom=0, left=0, right=0)
        _margens_celula(tbl_ft.cell(0, 1), top=60, bottom=0, left=0, right=0)

        p_left = tbl_ft.cell(0, 0).paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _aplicar_formatacao_paragrafo(p_left, doc.get("rodape", ""), tamanho_pt=9.0, cor_rgb=RGBColor(60, 60, 60))

        p_right = tbl_ft.cell(0, 1).paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _adicionar_campo_pagina_docx(p_right, tamanho_pt=9.0, cor_rgb=RGBColor(60, 60, 60))

    unid = config.unidade()
    ident = doc["identificacao"]

    # 1. Brasão das Armas Nacionais
    brasao = _brasao_bytes()
    if brasao:
        p_brasao = d.add_paragraph()
        p_brasao.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_brasao.paragraph_format.space_before = Pt(0)
        p_brasao.paragraph_format.space_after = Pt(2)
        run_b = p_brasao.add_run()
        run_b.add_picture(io.BytesIO(brasao), height=Cm(1.8))

    # 2. Cabeçalho Institucional
    for linha in unid.get("cabecalho", []):
        p_cab = d.add_paragraph()
        p_cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cab.paragraph_format.space_before = Pt(0)
        p_cab.paragraph_format.space_after = Pt(1)
        p_cab.paragraph_format.line_spacing = Pt(13)
        _aplicar_formatacao_paragrafo(p_cab, linha, tamanho_pt=11.5, negrito_padrao=True)

    # 3. Identificação: DIEx e NUP à esquerda, Local e Data à direita
    p_espaco = d.add_paragraph()
    p_espaco.paragraph_format.space_before = Pt(6)
    p_espaco.paragraph_format.space_after = Pt(4)

    tbl_id = d.add_table(rows=1, cols=2)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_id.autofit = False
    tbl_id.columns[0].width = Cm(10.5)
    tbl_id.columns[1].width = Cm(7.2)

    cell_esq = tbl_id.cell(0, 0)
    cell_dir = tbl_id.cell(0, 1)

    _definir_bordas_celula(cell_esq, top="none", bottom="none", left="none", right="none")
    _definir_bordas_celula(cell_dir, top="none", bottom="none", left="none", right="none")

    p_id_esq1 = cell_esq.paragraphs[0]
    p_id_esq1.paragraph_format.space_after = Pt(1)
    _aplicar_formatacao_paragrafo(p_id_esq1, ident.get("titulo", ""), tamanho_pt=12, negrito_padrao=True)

    p_id_esq2 = cell_esq.add_paragraph()
    p_id_esq2.paragraph_format.space_after = Pt(0)
    _aplicar_formatacao_paragrafo(p_id_esq2, f"NUP: {ident.get('nup', '')}", tamanho_pt=12, negrito_padrao=True)

    p_id_dir = cell_dir.paragraphs[0]
    p_id_dir.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_id_dir.paragraph_format.space_after = Pt(0)
    _aplicar_formatacao_paragrafo(p_id_dir, f"Rio de Janeiro-RJ, {ident.get('data', '')}", tamanho_pt=12)

    # Espaçamento pós identificação
    p_sep = d.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(4)
    p_sep.paragraph_format.space_after = Pt(2)

    # Do / Ao / Assunto
    for rotulo, val in (("Do", ident.get("do")),
                        ("Ao", ident.get("ao", "").removeprefix("Ao ")),
                        ("Assunto:", ident.get("assunto"))):
        p_meta = d.add_paragraph()
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(2)
        _aplicar_formatacao_paragrafo(p_meta, f"<b>{rotulo}</b> {val or ''}", tamanho_pt=12)

    # Anexos
    if doc.get("anexos"):
        p_anx_tit = d.add_paragraph()
        p_anx_tit.paragraph_format.space_before = Pt(4)
        p_anx_tit.paragraph_format.space_after = Pt(1)
        _aplicar_formatacao_paragrafo(p_anx_tit, "<b>Anexos:</b>", tamanho_pt=12)
        for k, a in enumerate(doc["anexos"], 1):
            p_anx = d.add_paragraph()
            p_anx.paragraph_format.left_indent = Cm(1.0)
            p_anx.paragraph_format.space_before = Pt(0)
            p_anx.paragraph_format.space_after = Pt(1)
            num = f"{k}) " if len(doc["anexos"]) > 1 else ""
            _aplicar_formatacao_paragrafo(p_anx, f"{num}{a}", tamanho_pt=11.5)

    p_sep2 = d.add_paragraph()
    p_sep2.paragraph_format.space_before = Pt(4)
    p_sep2.paragraph_format.space_after = Pt(4)

    # 4. Parágrafos do Texto
    tab_c = doc.get("tabela_credito")
    for p_texto in doc.get("paragrafos", []):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        _aplicar_formatacao_paragrafo(p, p_texto, tamanho_pt=12)

        if tab_c and ("Fonte de Recurso" in p_texto or "crédito orçamentário" in p_texto):
            p_sep_c = d.add_paragraph()
            p_sep_c.paragraph_format.space_before = Pt(2)
            p_sep_c.paragraph_format.space_after = Pt(2)

            tbl_c = d.add_table(rows=2 if not tab_c.get("detalhamento") else 3, cols=8)
            tbl_c.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_c.autofit = False

            larguras_c = [Cm(2.8), Cm(1.7), Cm(1.7), Cm(1.8), Cm(3.2), Cm(1.6), Cm(1.6), Cm(3.3)]
            for col_i, w in enumerate(larguras_c):
                tbl_c.columns[col_i].width = w

            # Cabeçalho da tabela de crédito
            titulos_c = ("NOTA DE CRÉDITO", "UG EMIT", "UG EXEC", "ND", "PLANO INTERNO (PI)", "PTRES", "UGR", "VALOR SOLICITADO")
            for col_i, h_tit in enumerate(titulos_c):
                cell_h = tbl_c.cell(0, col_i)
                _sombrear_celula(cell_h, "EEF2F6")
                _definir_bordas_celula(cell_h, cor="718096", sz="4")
                _margens_celula(cell_h, 50, 50, 40, 40)
                p_ch = cell_h.paragraphs[0]
                p_ch.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ch.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p_ch, h_tit, tamanho_pt=8.5, negrito_padrao=True)

            # Linha de dados
            vals_c = [
                (tab_c.get("nc", "—"), WD_ALIGN_PARAGRAPH.CENTER, True),
                (tab_c.get("emitente", "—"), WD_ALIGN_PARAGRAPH.CENTER, False),
                (tab_c.get("ug", "—"), WD_ALIGN_PARAGRAPH.CENTER, False),
                (tab_c.get("nd", "—"), WD_ALIGN_PARAGRAPH.CENTER, True),
                (tab_c.get("pi", "—"), WD_ALIGN_PARAGRAPH.CENTER, False),
                (tab_c.get("ptres", "—"), WD_ALIGN_PARAGRAPH.CENTER, False),
                (tab_c.get("ugr", "—"), WD_ALIGN_PARAGRAPH.CENTER, False),
                (tab_c.get("valor", "—"), WD_ALIGN_PARAGRAPH.RIGHT, True),
            ]
            for col_i, (v_txt, v_align, v_neg) in enumerate(vals_c):
                cell_d = tbl_c.cell(1, col_i)
                _definir_bordas_celula(cell_d, cor="718096", sz="4")
                _margens_celula(cell_d, 50, 50, 40, 40)
                p_cd = cell_d.paragraphs[0]
                p_cd.alignment = v_align
                p_cd.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p_cd, v_txt, tamanho_pt=9.5, negrito_padrao=v_neg)

            # Linha de detalhamento (se houver)
            if tab_c.get("detalhamento"):
                cell_det = tbl_c.cell(2, 0)
                for ck in range(1, 8):
                    cell_det.merge(tbl_c.cell(2, ck))
                _sombrear_celula(cell_det, "F8FAFC")
                _definir_bordas_celula(cell_det, cor="718096", sz="4")
                _margens_celula(cell_det, 50, 50, 60, 60)
                p_det = cell_det.paragraphs[0]
                p_det.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p_det, f"<b>Detalhamento de Crédito:</b> {tab_c['detalhamento']}", tamanho_pt=8.5)

            p_pos_c = d.add_paragraph()
            p_pos_c.paragraph_format.space_before = Pt(4)
            p_pos_c.paragraph_format.space_after = Pt(6)
            tab_c = None

    # 5. Quadro de Condições (se houver, ex.: Requisição)
    quadro = doc.get("quadro")
    if quadro and isinstance(quadro, (list, tuple)) and len(quadro) > 0:
        tbl_q = d.add_table(rows=len(quadro), cols=2)
        tbl_q.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_q.autofit = False
        tbl_q.columns[0].width = Cm(5.5)
        tbl_q.columns[1].width = Cm(12.2)

        for i, item in enumerate(quadro):
            c0 = tbl_q.cell(i, 0)
            c1 = tbl_q.cell(i, 1)
            _definir_bordas_celula(c0, cor="888888", sz="4")
            _definir_bordas_celula(c1, cor="888888", sz="4")
            _sombrear_celula(c0, "F8F8F8")
            _margens_celula(c0, 60, 60, 100, 100)
            _margens_celula(c1, 60, 60, 100, 100)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p0, item[0] if isinstance(item, (list, tuple)) else str(item),
                                          tamanho_pt=10.5, negrito_padrao=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p1, item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else "",
                                          tamanho_pt=10.5)

        p_pos_q = d.add_paragraph()
        p_pos_q.paragraph_format.space_before = Pt(4)
        p_pos_q.paragraph_format.space_after = Pt(4)

    # 6. Tabelas de Itens (Requisição e Reforço) ou Anulação
    tabelas = doc.get("tabelas") or []
    tipo_doc = doc.get("tipo", "REQUISICAO")

    # Padronização: A tabela de itens fica sempre na segunda página
    if tabelas:
        d.add_page_break()

    if tipo_doc == "ANULACAO":
        for tab in tabelas:
            linhas = tab.get("linhas", [])
            tbl = d.add_table(rows=len(linhas) + 3, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            tbl.columns[0].width = Cm(11.5)
            tbl.columns[1].width = Cm(6.2)

            # Cabeçalho 1: Fornecedor (mesclado)
            c_forn = tbl.cell(0, 0)
            c_forn.merge(tbl.cell(0, 1))
            _sombrear_celula(c_forn, "E8E8E8")
            _definir_bordas_celula(c_forn)
            _margens_celula(c_forn, 80, 80, 100, 100)
            p_forn = c_forn.paragraphs[0]
            p_forn.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_forn, tab.get("cabecalho", ""), tamanho_pt=10.5, negrito_padrao=True)

            # Cabeçalho 2: Colunas
            for col_idx, h_tit in enumerate(("NOTA DE EMPENHO", "VALOR A ANULAR")):
                c = tbl.cell(1, col_idx)
                _sombrear_celula(c, "F2F2F2")
                _definir_bordas_celula(c)
                _margens_celula(c, 60, 60, 100, 100)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p, h_tit, tamanho_pt=10, negrito_padrao=True)

            # Linhas
            for r_idx, linha in enumerate(linhas, 2):
                c0 = tbl.cell(r_idx, 0)
                c1 = tbl.cell(r_idx, 1)
                _definir_bordas_celula(c0)
                _definir_bordas_celula(c1)
                _margens_celula(c0, 50, 50, 80, 80)
                _margens_celula(c1, 50, 50, 80, 80)

                p0 = c0.paragraphs[0]
                p0.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p0, f"<b>{linha.get('ne', '')}</b> · {linha.get('detalhe', '')}", tamanho_pt=10)

                p1 = c1.paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p1.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p1, linha.get("valor", ""), tamanho_pt=10)

            # Total do Fornecedor
            r_tot = len(linhas) + 2
            c_tot_rot = tbl.cell(r_tot, 0)
            c_tot_val = tbl.cell(r_tot, 1)
            _sombrear_celula(c_tot_rot, "F4F4F4")
            _sombrear_celula(c_tot_val, "F4F4F4")
            _definir_bordas_celula(c_tot_rot)
            _definir_bordas_celula(c_tot_val)
            _margens_celula(c_tot_rot, 70, 70, 100, 100)
            _margens_celula(c_tot_val, 70, 70, 100, 100)

            p_tr = c_tot_rot.paragraphs[0]
            p_tr.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_tr, "TOTAL DO FORNECEDOR", tamanho_pt=10, negrito_padrao=True)

            p_tv = c_tot_val.paragraphs[0]
            p_tv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_tv.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_tv, tab.get("total", ""), tamanho_pt=10, negrito_padrao=True)

            p_fim_t = d.add_paragraph()
            p_fim_t.paragraph_format.space_after = Pt(4)

    else:
        # REQUISICAO e REFORCO (Tabela completa com 7 colunas)
        # Colunas: ITEM, SI, DESCRIÇÃO, UND, QTDE, VALOR UNIT, VALOR TOTAL
        larguras = [Cm(1.4), Cm(1.2), Cm(6.3), Cm(1.4), Cm(2.2), Cm(2.5), Cm(2.7)]

        for tab in tabelas:
            linhas = tab.get("linhas", [])
            tbl = d.add_table(rows=len(linhas) + 3, cols=7)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False

            for c_idx, w in enumerate(larguras):
                tbl.columns[c_idx].width = w

            # Linha 0: Cabeçalho com Fornecedor (mesclado nas 7 colunas)
            c_forn = tbl.cell(0, 0)
            for c_k in range(1, 7):
                c_forn.merge(tbl.cell(0, c_k))
            _sombrear_celula(c_forn, "E8E8E8")
            _definir_bordas_celula(c_forn)
            _margens_celula(c_forn, 80, 80, 100, 100)
            p_forn = c_forn.paragraphs[0]
            p_forn.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_forn, tab.get("cabecalho", ""), tamanho_pt=10.5, negrito_padrao=True)

            # Linha 1: Títulos das colunas
            titulos = ("ITEM", "SI", "DESCRIÇÃO", "UND", "QTDE", "VALOR UNIT", "VALOR TOTAL")
            for col_idx, h_tit in enumerate(titulos):
                c = tbl.cell(1, col_idx)
                _sombrear_celula(c, "F2F2F2")
                _definir_bordas_celula(c)
                _margens_celula(c, 60, 60, 60, 60)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                _aplicar_formatacao_paragrafo(p, h_tit, tamanho_pt=9.5, negrito_padrao=True)

            # Linhas de Itens
            for r_idx, linha in enumerate(linhas, 2):
                valores = [
                    (linha.get("item", ""), WD_ALIGN_PARAGRAPH.CENTER, False),
                    (linha.get("si", ""), WD_ALIGN_PARAGRAPH.CENTER, False),
                    (linha.get("descricao", ""), WD_ALIGN_PARAGRAPH.LEFT, False),
                    (linha.get("und", ""), WD_ALIGN_PARAGRAPH.CENTER, False),
                    (linha.get("qtde", ""), WD_ALIGN_PARAGRAPH.RIGHT, False),
                    (linha.get("valor_unit", ""), WD_ALIGN_PARAGRAPH.RIGHT, False),
                    (linha.get("total", ""), WD_ALIGN_PARAGRAPH.RIGHT, False),
                ]
                for col_idx, (val, alin, neg) in enumerate(valores):
                    c = tbl.cell(r_idx, col_idx)
                    _definir_bordas_celula(c)
                    _margens_celula(c, 50, 50, 60, 60)
                    p = c.paragraphs[0]
                    p.alignment = alin
                    p.paragraph_format.space_after = Pt(0)
                    _aplicar_formatacao_paragrafo(p, str(val), tamanho_pt=9.5, negrito_padrao=neg)

            # Linha Total do Fornecedor (mescla 0 a 5, valor na coluna 6)
            r_tot = len(linhas) + 2
            c_tot_rot = tbl.cell(r_tot, 0)
            for c_k in range(1, 6):
                c_tot_rot.merge(tbl.cell(r_tot, c_k))
            c_tot_val = tbl.cell(r_tot, 6)

            _sombrear_celula(c_tot_rot, "F4F4F4")
            _sombrear_celula(c_tot_val, "F4F4F4")
            _definir_bordas_celula(c_tot_rot)
            _definir_bordas_celula(c_tot_val)
            _margens_celula(c_tot_rot, 70, 70, 100, 100)
            _margens_celula(c_tot_val, 70, 70, 100, 100)

            p_tr = c_tot_rot.paragraphs[0]
            p_tr.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_tr, "TOTAL DO FORNECEDOR", tamanho_pt=10, negrito_padrao=True)

            p_tv = c_tot_val.paragraphs[0]
            p_tv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_tv.paragraph_format.space_after = Pt(0)
            _aplicar_formatacao_paragrafo(p_tv, tab.get("total", ""), tamanho_pt=10, negrito_padrao=True)

            p_fim_t = d.add_paragraph()
            p_fim_t.paragraph_format.space_after = Pt(4)

    # Total Geral (se presente)
    if doc.get("total_geral"):
        p_tot_g = d.add_paragraph()
        p_tot_g.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_tot_g.paragraph_format.space_before = Pt(4)
        p_tot_g.paragraph_format.space_after = Pt(8)
        _aplicar_formatacao_paragrafo(p_tot_g, f"<b>VALOR TOTAL:</b> {doc['total_geral']}", tamanho_pt=11.5)

    # 7. Assinatura do Requisitante
    p_pre_ass = d.add_paragraph()
    p_pre_ass.paragraph_format.space_before = Pt(14)
    p_pre_ass.paragraph_format.space_after = Pt(0)

    ass = doc.get("assinatura", {})
    if doc.get("modo_assinatura") == "DIGITAL":
        p_dig = d.add_paragraph()
        p_dig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dig.paragraph_format.space_after = Pt(2)
        _aplicar_formatacao_paragrafo(p_dig, "[ Espaço para Assinatura Digital do Requisitante ]",
                                      tamanho_pt=8, italico=True, cor_rgb=RGBColor(100, 100, 100))

    p_ass_nome = d.add_paragraph()
    p_ass_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass_nome.paragraph_format.space_before = Pt(12)
    p_ass_nome.paragraph_format.space_after = Pt(1)
    _aplicar_formatacao_paragrafo(p_ass_nome, ass.get("nome", ""), tamanho_pt=12, negrito_padrao=True)

    p_ass_func = d.add_paragraph()
    p_ass_func.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass_func.paragraph_format.space_after = Pt(14)
    _aplicar_formatacao_paragrafo(p_ass_func, ass.get("funcao", ""), tamanho_pt=11)

    # 8. Quadro de Despacho (Fiscal Administrativo e Ordenador de Despesas)
    desp = doc.get("despacho")
    if desp:
        # Padronização: O despacho fica sempre na próxima página
        d.add_page_break()
        p_desp_tit = d.add_paragraph()
        p_desp_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_desp_tit.paragraph_format.space_before = Pt(10)
        p_desp_tit.paragraph_format.space_after = Pt(4)
        _aplicar_formatacao_paragrafo(p_desp_tit, "<b>DESPACHO</b>", tamanho_pt=12)

        tbl_desp = d.add_table(rows=1, cols=2)
        tbl_desp.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_desp.autofit = False
        tbl_desp.columns[0].width = Cm(8.8)
        tbl_desp.columns[1].width = Cm(8.8)

        cf = tbl_desp.cell(0, 0)
        co = tbl_desp.cell(0, 1)

        _definir_bordas_celula(cf, top="single", bottom="single", left="single", right="single", cor="000000", sz="4")
        _definir_bordas_celula(co, top="single", bottom="single", left="single", right="single", cor="000000", sz="4")
        _margens_celula(cf, 120, 120, 140, 140)
        _margens_celula(co, 120, 120, 140, 140)

        # Coluna Fiscal
        df = desp.get("fiscal", {})
        pf_tit = cf.paragraphs[0]
        pf_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf_tit.paragraph_format.space_after = Pt(6)
        _aplicar_formatacao_paragrafo(pf_tit, f"<u><b>{df.get('titulo', '')}</b></u>", tamanho_pt=10.5)

        for t in df.get("texto", []):
            pt = cf.add_paragraph()
            pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pt.paragraph_format.space_after = Pt(4)
            _aplicar_formatacao_paragrafo(pt, t, tamanho_pt=10.5)

        pf_data = cf.add_paragraph()
        pf_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf_data.paragraph_format.space_before = Pt(12)
        pf_data.paragraph_format.space_after = Pt(18)
        _aplicar_formatacao_paragrafo(pf_data, "Rio de Janeiro-RJ, ___ / ___ / ___.", tamanho_pt=10)

        pf_nome = cf.add_paragraph()
        pf_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf_nome.paragraph_format.space_after = Pt(1)
        _aplicar_formatacao_paragrafo(pf_nome, df.get("nome", ""), tamanho_pt=10.5, negrito_padrao=True)

        pf_func = cf.add_paragraph()
        pf_func.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf_func.paragraph_format.space_after = Pt(0)
        _aplicar_formatacao_paragrafo(pf_func, df.get("funcao", ""), tamanho_pt=10)

        # Coluna OD
        do = desp.get("od", {})
        po_tit = co.paragraphs[0]
        po_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        po_tit.paragraph_format.space_after = Pt(6)
        _aplicar_formatacao_paragrafo(po_tit, f"<u><b>{do.get('titulo', '')}</b></u>", tamanho_pt=10.5)

        for t in do.get("texto", []):
            pt = co.add_paragraph()
            pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pt.paragraph_format.space_after = Pt(4)
            _aplicar_formatacao_paragrafo(pt, t, tamanho_pt=10.5)

        po_data = co.add_paragraph()
        po_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
        po_data.paragraph_format.space_before = Pt(12)
        po_data.paragraph_format.space_after = Pt(18)
        _aplicar_formatacao_paragrafo(po_data, "Rio de Janeiro-RJ, ___ / ___ / ___.", tamanho_pt=10)

        po_nome = co.add_paragraph()
        po_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
        po_nome.paragraph_format.space_after = Pt(1)
        _aplicar_formatacao_paragrafo(po_nome, do.get("nome", ""), tamanho_pt=10.5, negrito_padrao=True)

        po_func = co.add_paragraph()
        po_func.alignment = WD_ALIGN_PARAGRAPH.CENTER
        po_func.paragraph_format.space_after = Pt(0)
        _aplicar_formatacao_paragrafo(po_func, do.get("funcao", ""), tamanho_pt=10)

    # -------------------------------------------------------------------------
    # Anexos incorporados ao documento Word (NC, NE, etc.)
    # -------------------------------------------------------------------------
    if anexos:
        import pymupdf
        for rotulo, arq_anexo in anexos:
            p_arq = Path(arq_anexo) if arq_anexo else None
            if not p_arq or not p_arq.is_file():
                continue
            try:
                with pymupdf.open(str(p_arq)) as doc_pdf:
                    total_pags = len(doc_pdf)
                    for idx_pag, pag in enumerate(doc_pdf):
                        d.add_page_break()
                        p_tit = d.add_paragraph()
                        p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_tit.paragraph_format.space_before = Pt(4)
                        p_tit.paragraph_format.space_after = Pt(8)
                        pag_info = f" (Pág. {idx_pag + 1}/{total_pags})" if total_pags > 1 else ""
                        _aplicar_formatacao_paragrafo(p_tit, f"<b>ANEXO: {rotulo.upper()}{pag_info}</b>", tamanho_pt=11)
                        pix = pag.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        d.add_picture(io.BytesIO(img_bytes), width=Cm(17.5))
            except Exception:
                pass

    d.save(str(destino))
    return destino
