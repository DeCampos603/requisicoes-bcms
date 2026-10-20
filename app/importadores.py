# -*- coding: utf-8 -*-
"""Importadores: transformam as fontes do repositório em índices JSON em dados/.

  nes       PDFs de Notas de Empenho (EMPENHOS/Empenhos/<UG>/)         -> dados/nes.json
  ncs       PDFs de Notas de Crédito (EMPENHOS/Notas de Crédito/<UG>/) -> dados/ncs.json
  credito   saldos por NC do Dashboard-Credito-BCMS (site/index.html)  -> dados/credito.json
  pregoes   Consolidado_Pregoes.xlsx (base do dashboard de capacidade) -> dados/pregoes.json
  uasgs     nomes das UASG (NC emitentes + consolidado)                -> dados/uasgs.json

Uso:  py -3 -m app.importadores todos
      py -3 -m app.importadores pregoes --xlsx "caminho\\Consolidado_Pregoes.xlsx"
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import threading
import time
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Callable

from . import config
from .dinheiro import ZERO, ValorInvalido, parse_decimal
from .leitor_nc import ler_nc
from .leitor_ne import completar_registro, ler_ne

Progresso = Callable[[str], None]


def _versao(registro: dict) -> int:
    """Versão do documento como número ('009' < '010'; como texto, '999' > '1000')."""
    try:
        return int(registro.get("versao") or 0)
    except (TypeError, ValueError):
        return 0


def _gravar(nome: str, dados: dict) -> Path:
    config.DADOS.mkdir(parents=True, exist_ok=True)
    destino = config.DADOS / nome
    with config.ARQUIVO_LOCK:
        tmp = destino.with_suffix(f".{os.getpid()}.{threading.get_ident()}.tmp")
        tmp.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(destino)  # troca atômica: nunca deixa índice pela metade
    return destino


def _agora() -> str:
    return datetime.now().strftime("%d/%m/%Y %H:%M")


# ----------------------------------------------------------------------------- NE
def _pastas_ne() -> dict[str, list[Path]]:
    pastas: dict[str, list[Path]] = {}
    for ug in config.NE_DIRS:
        pastas[ug] = [config.NE_DIRS[ug], config.UPLOAD_NE_DIRS[ug], config.DADOS_NE_DIRS[ug]]
    return pastas


def importar_nes(progresso: Progresso = print) -> dict:
    todas_pastas = [p for pastas in _pastas_ne().values() for p in pastas if p.is_dir()]
    if not todas_pastas:
        raise FileNotFoundError("Pastas de empenhos (NE) não encontradas em EMPENHOS nem em dados/ne.")
    nes: dict[str, dict] = {}
    falhas: list[dict] = []
    import shutil
    for ug, pastas in _pastas_ne().items():
        destino_local = config.DADOS_NE_DIRS[ug]
        destino_local.mkdir(parents=True, exist_ok=True)
        for pasta in dict.fromkeys(pastas):  # sem repetir quando o upload grava na pasta oficial
            if not pasta.is_dir():
                continue
            for pdf in sorted(pasta.glob("*.pdf")):
                copia_local = destino_local / pdf.name
                if pasta != destino_local and not copia_local.is_file():
                    try:
                        shutil.copy2(pdf, copia_local)
                    except Exception:
                        pass
                try:
                    ne = ler_ne(pdf)
                except Exception as e:  # PDF corrompido não derruba a importação
                    falhas.append({"arquivo": pdf.name, "erro": str(e)})
                    continue
                if not ne.get("chave"):
                    falhas.append({"arquivo": pdf.name, "erro": "NE sem número identificável"})
                    continue
                ne["ug_pasta"] = ug
                ne["caminho"] = str(pdf)
                completar_registro(ne)
                chave = f"{ug}:{ne['chave']}"
                antiga = nes.get(chave)
                # se a mesma NE existe mais de uma vez, vale a versão numericamente maior
                if antiga is None or _versao(ne) >= _versao(antiga):
                    nes[chave] = ne
        progresso(f"  NE {ug}: {sum(1 for k in nes if k.startswith(ug + ':'))} lidas")
    saida = {"gerado_em": _agora(), "total": len(nes), "falhas": falhas, "nes": nes}
    _gravar("nes.json", saida)
    return {"nes": len(nes), "falhas": len(falhas)}


# ----------------------------------------------------------------------------- NC
def importar_ncs(progresso: Progresso = print) -> dict:
    todas_pastas = [p for ug in config.NC_DIRS for p in (config.NC_DIRS[ug], config.UPLOAD_NC_DIRS[ug], config.DADOS_NC_DIRS[ug]) if p.is_dir()]
    if not todas_pastas:
        raise FileNotFoundError("Pastas de notas de crédito (NC) não encontradas em EMPENHOS nem em dados/nc.")
    ncs: dict[str, dict] = {}
    falhas: list[dict] = []
    uasgs: dict[str, str] = {}
    import shutil
    for ug in config.NC_DIRS:
        destino_local = config.DADOS_NC_DIRS[ug]
        destino_local.mkdir(parents=True, exist_ok=True)
        for pasta in dict.fromkeys((config.NC_DIRS[ug], config.UPLOAD_NC_DIRS[ug], config.DADOS_NC_DIRS[ug])):
            if not pasta.is_dir():
                continue
            for pdf in sorted(pasta.glob("*.pdf")):
                copia_local = destino_local / pdf.name
                if pasta != destino_local and not copia_local.is_file():
                    try:
                        shutil.copy2(pdf, copia_local)
                    except Exception:
                        pass
                try:
                    nc = ler_nc(pdf)
                except Exception as e:
                    falhas.append({"arquivo": pdf.name, "erro": str(e)})
                    continue
                if not nc.get("chave") or not nc.get("ug_emitente"):
                    falhas.append({"arquivo": pdf.name, "erro": "NC sem número ou emitente identificável"})
                    continue
                nc["ug_pasta"] = ug
                nc["caminho"] = str(pdf)
                if nc.get("nome_emitente"):
                    uasgs[nc["ug_emitente"]] = nc["nome_emitente"]
                for it in nc["itens"]:
                    uasgs.setdefault(it["ug_favorecida"], it["nome_favorecida"])
                chave = f"{nc['ug_emitente']}:{nc['chave']}"
                antiga = ncs.get(chave)
                if antiga is None or _versao(nc) >= _versao(antiga):
                    ncs[chave] = nc
        progresso(f"  NC {ug}: {sum(1 for v in ncs.values() if v['ug_pasta'] == ug)} lidas")
    _gravar("ncs.json", {"gerado_em": _agora(), "total": len(ncs), "falhas": falhas, "ncs": ncs})
    _mesclar_uasgs(uasgs, "nc")
    return {"ncs": len(ncs), "falhas": len(falhas)}


# ----------------------------------------------------------------------------- crédito (dashboard)
def _bloco_js(texto: str, nome: str):
    m = re.search(rf"(?:const|var|let)\s+{nome}\s*=\s*", texto)
    if not m:
        raise ValueError(f"Bloco {nome} não encontrado no HTML do dashboard (layout mudou?).")
    return json.JSONDecoder().raw_decode(texto[m.end():])[0]


def importar_credito(progresso: Progresso = print) -> dict:
    caminho = config.CREDITO_HTML
    if not caminho.is_file():
        raise FileNotFoundError(f"Dashboard de crédito não encontrado: {caminho}")
    html = caminho.read_text(encoding="utf-8")
    hist = _bloco_js(html, "HISTDATA")
    m = re.search(r"atualizado em (\d{2}/\d{2}/\d{4}) às (\d{2}:\d{2})", html)
    posicao = f"{m.group(1)} {m.group(2)}" if m else None
    ncs: dict[str, dict] = {}
    for it in hist["items"]:
        if it["fav_cod"] not in ("160329", "167329"):
            continue
        chave = f"{it['emit_cod']}:{it['nc'][-12:]}"  # UG emitente : AAAANC999999
        cel = it.get("itens") or [{"acao": it["acao"], "pi": it["pi"], "nd": it["nd"],
                                    "prov": it["prov"], "emp": it["emp"], "cred": it["cred"]}]
        reg = ncs.setdefault(chave, {
            "emit": it["emit_cod"], "chave": it["nc"][-12:], "dia": it["dia"], "op": it["op"],
            "obj": it.get("obj"), "favorecidas": {},
        })
        reg["favorecidas"][it["fav_cod"]] = {
            "status": it.get("status"),
            "prov": it["prov"], "emp": it["emp"], "cred": it["cred"],
            "celulas": [{"ptres": c["acao"], "pi": c["pi"], "nd": c["nd"], "prov": c["prov"],
                         "emp": c["emp"], "cred": c["cred"]} for c in cel],
        }
    _gravar("credito.json", {"posicao": posicao, "origem": str(caminho), "importado_em": _agora(),
                             "total": len(ncs), "ncs": ncs})
    progresso(f"  Crédito: {len(ncs)} NC do BCMS (posição {posicao})")
    return {"ncs": len(ncs), "posicao": posicao}


# ----------------------------------------------------------------------------- pregões
def _num(v) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return parse_decimal(v)
    except ValorInvalido:
        return None


def _data_br(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    m = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", s)
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else None


def importar_pregoes(xlsx: Path | None = None, progresso: Progresso = print) -> dict:
    from openpyxl import load_workbook

    xlsx = Path(xlsx) if xlsx else config.PREGOES_XLSX
    if not xlsx.is_file():
        raise FileNotFoundError(f"Consolidado de pregões não encontrado: {xlsx}")
    uasg_propria = config.unidade()["uasg_propria"]
    papel_de = {c["consolidado"]: c["id"] for c in config.catalogos()["tipos_unidade"]}
    progresso(f"  Lendo {xlsx.name} (pode levar cerca de 1 minuto)...")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = ws.iter_rows(values_only=True)
    cab = [str(c or "").strip() for c in next(linhas)]
    ix = {c: i for i, c in enumerate(cab)}
    exigidas = ["UASG_Gerenciadora", "Pregão", "Nr Item", "UASG", "Tipo UASG", "Val. Unitário", "Qtd. Saldo"]
    faltam = [c for c in exigidas if c not in ix]
    if faltam:
        raise ValueError(f"Layout do consolidado mudou: colunas ausentes {faltam}")

    def g(r, col):
        i = ix.get(col)
        return r[i] if i is not None and i < len(r) else None

    pregoes: dict[str, dict] = {}
    nomes_uasg: dict[str, str] = {}
    datas_coleta: list[str] = []
    lidas = repetidas = 0
    for r in linhas:
        lidas += 1
        uasg_txt = str(g(r, "UASG") or "")
        m_u = re.match(r"(\d{6})\s*-\s*(.+)", uasg_txt)
        if m_u:
            nomes_uasg.setdefault(m_u.group(1), m_u.group(2).strip())
        if not uasg_txt.startswith(uasg_propria):
            continue
        ugg = str(g(r, "UASG_Gerenciadora") or "").strip()
        numero = str(g(r, "Pregão") or "").replace("_", "/")
        papel = papel_de.get(str(g(r, "Tipo UASG") or "").strip())
        if not (ugg and numero and papel):
            continue
        chave = f"{ugg}:{numero}:{papel}"
        reg = pregoes.setdefault(chave, {
            "id": chave, "ugg": ugg, "numero": numero, "modalidade": "PREGAO",
            "tipo_unidade": papel, "itens": [],
        })
        forn = str(g(r, "Fornecedor") or "")
        m_f = re.match(r"([\d./A-Z-]+?)\s+-\s+(.+)", forn)
        d = _data_br(g(r, "Data_Coleta"))
        if d:
            datas_coleta.append(d)
        linha_item = {
            "item": g(r, "Nr Item"),
            "ata": str(g(r, "Nr Ata") or "").strip() or None,
            "descricao": str(g(r, "Descrição detalhada") or "").strip(),
            "catmat": str(g(r, "CATMAT_CATSER_Codigo") or "").strip() or None,
            "catmat_nome": str(g(r, "CATMAT_CATSER_Nome") or "").strip() or None,
            "catalogo": str(g(r, "Catalogo_Tipo") or "").strip() or None,
            "categoria": str(g(r, "Categoria_Geral") or "").strip() or None,
            "nd_sugerida": str(g(r, "ND_Sugerida") or "").strip() or None,
            "subitem_sugerido": str(g(r, "Subitem_Sugerido") or "").strip() or None,
            "fornecedor_cnpj": m_f.group(1) if m_f else None,
            "fornecedor_nome": m_f.group(2).strip() if m_f else (forn.strip() or None),
            "vigencia_ini": _data_br(g(r, "Início Vig Ata")),
            "vigencia_fim": _data_br(g(r, "Fim Vig Ata")),
            "status": str(g(r, "Status_Ata") or "").strip() or None,
            "valor_unitario": str(_num(g(r, "Val. Unitário"))) if _num(g(r, "Val. Unitário")) is not None else None,
            "qtd_homologada": str(_num(g(r, "Qtd Homologada"))) if _num(g(r, "Qtd Homologada")) is not None else None,
            "qtd_autorizada": str(_num(g(r, "Qtd. Autorizada"))) if _num(g(r, "Qtd. Autorizada")) is not None else None,
            "qtd_saldo": str(_num(g(r, "Qtd. Saldo"))) if _num(g(r, "Qtd. Saldo")) is not None else None,
            "link": str(g(r, "Link_Item") or "").strip() or None,
        }
        if linha_item in reg["itens"]:  # a planilha repete linhas idênticas (mesmo item/ata/fornecedor/saldo): conta uma vez só
            repetidas += 1
            continue
        reg["itens"].append(linha_item)
    wb.close()
    from .base import sintetizar_objeto_pregao
    for reg in pregoes.values():
        reg["itens"].sort(key=lambda i: (i["item"] if isinstance(i["item"], int) else 10**9))
        reg["n_itens"] = len(reg["itens"])
        fins = [i["vigencia_fim"] for i in reg["itens"] if i["vigencia_fim"]]
        reg["vigencia_fim_min"] = min(fins) if fins else None
        reg["vigencia_fim_max"] = max(fins) if fins else None
        reg["fornecedores"] = sorted({i["fornecedor_nome"] for i in reg["itens"] if i["fornecedor_nome"]})
        obj, cats = sintetizar_objeto_pregao(reg)
        reg["objeto"] = obj
        reg["categorias"] = cats
    posicao = max(datas_coleta) if datas_coleta else None
    _gravar("pregoes.json", {
        "posicao": posicao, "origem": xlsx.name, "importado_em": _agora(),
        "linhas_lidas": lidas, "total": len(pregoes),
        "pregoes": dict(sorted(pregoes.items())),
    })
    _mesclar_uasgs(nomes_uasg, "consolidado")
    progresso(f"  Pregões: {len(pregoes)} (pregão × UASG gerenciadora × papel do BCMS); posição {posicao}"
              + (f"; {repetidas} linha(s) idêntica(s) repetida(s) na planilha ignorada(s)" if repetidas else ""))
    return {"pregoes": len(pregoes), "posicao": posicao, "linhas_repetidas": repetidas}


# ----------------------------------------------------------------------------- UASG
def _mesclar_uasgs(novos: dict[str, str], origem: str) -> None:
    arq = config.DADOS / "uasgs.json"
    atual = json.loads(arq.read_text(encoding="utf-8")) if arq.is_file() else {"uasgs": {}}
    for cod, nome in novos.items():
        reg = atual["uasgs"].setdefault(cod, {})
        # nome completo (NC) tem prioridade sobre a sigla do consolidado
        if origem == "nc":
            reg["nome"] = nome
        else:
            reg.setdefault("sigla", nome)
    atual["atualizado_em"] = _agora()
    _gravar("uasgs.json", atual)


# ----------------------------------------------------------------------------- CLI
def importar_pca_etapa(progresso: Progresso = print) -> dict:
    from . import pca
    return pca.importar_pca(progresso)


def importar_catalogo_etapa(progresso: Progresso = print) -> dict:
    from . import pca
    return pca.importar_catalogo(progresso)


ETAPAS = {"nes": importar_nes, "ncs": importar_ncs, "credito": importar_credito, "pregoes": importar_pregoes,
          "pca": importar_pca_etapa, "catalogo": importar_catalogo_etapa}


def atualizar_tudo(progresso: Progresso = print, xlsx: Path | None = None) -> dict:
    resumo: dict = {}
    # pca e catalogo vêm da API pública (rede): a falha deles não impede o restante
    for nome in ("ncs", "nes", "credito", "pregoes", "pca", "catalogo"):
        t0 = time.time()
        progresso(f"[{nome}]")
        try:
            resumo[nome] = importar_pregoes(xlsx, progresso) if nome == "pregoes" else ETAPAS[nome](progresso)
        except Exception as e:  # uma fonte ausente não impede as demais
            resumo[nome] = {"erro": f"{type(e).__name__}: {e}"}
            progresso(f"  FALHOU: {e}")
        resumo[nome]["segundos"] = round(time.time() - t0, 1)
    return resumo


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("etapa", choices=["todos", *ETAPAS])
    ap.add_argument("--xlsx", help="Consolidado_Pregoes.xlsx alternativo")
    args = ap.parse_args(argv)
    xlsx = Path(args.xlsx) if args.xlsx else None
    if args.etapa == "todos":
        resumo = atualizar_tudo(xlsx=xlsx)
    elif args.etapa == "pregoes":
        resumo = {"pregoes": importar_pregoes(xlsx)}
    else:
        resumo = {args.etapa: ETAPAS[args.etapa]()}
    print(json.dumps(resumo, ensure_ascii=False, indent=1))
    return 1 if any("erro" in v for v in resumo.values()) else 0


if __name__ == "__main__":
    raise SystemExit(main())
